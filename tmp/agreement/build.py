from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path('/Users/levpolshkov/kami/output/documents/Kami_Aesthetics_NP_Agreement_December_2026.docx')
doc=Document()
s=doc.sections[0]
s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
s.page_width=Inches(8.5); s.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',19),('Heading 1',13),('Heading 2',10.5)]:
 st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True; st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(6)
 st.paragraph_format.keep_with_next=True
h=s.header.paragraphs[0]; h.text='KAMI AESTHETICS  |  APRN AGREEMENT'; h.style='Normal'; h.runs[0].font.size=Pt(8)
f=s.footer.paragraphs[0]; f.alignment=2
r=f.add_run('Kami Aesthetics  |  Page '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
doc.core_properties.title='Kami Aesthetics APRN Independent Contractor Agreement'
doc.core_properties.subject='Compensation and services effective December 1 2026'
doc.core_properties.author='Kami Aesthetics'
text='''
# Kami Aesthetics Independent Contractor Agreement
Advanced Practice Registered Nurse Medical Aesthetics Services
Aventura Aesthetics Group LLC d/b/a Kami Aesthetics
2999 NE 191st Street, Suite 906, Aventura, Florida 33180
Effective Date: December 1, 2026
This Agreement is between Aventura Aesthetics Group LLC d/b/a Kami Aesthetics (Company) and Valeriia Tiertyshnikova, APRN (Contractor). It governs Contractor's medical aesthetic services, compensation, professional responsibilities, and assistance with access to medical products. The parties acknowledge Company's investment in training, Contractor's increased clinical experience, and Contractor's purchasing relationships that provide Company access to favorable product pricing.
Contractor Florida APRN license number: 11043386. Contractor shall verify current credentials before services begin under this Agreement. Contractor's current address and notice email shall be entered on the signature page. The designated Medical Director is Dr. Paul M. Goldberg, MD, or a properly qualified replacement identified to Contractor in writing.
## 1 Independent contractor relationship
1.1 Status. The parties intend an independent contractor relationship, not an employment, partnership, or joint venture relationship. Actual working arrangements must comply with applicable classification law; this designation does not waive statutory rights or determine legal status by itself. Subject to applicable law, Contractor is responsible for her taxes, receives no employee benefits, and shall provide a completed Form W-9. Company shall issue required tax reporting forms.
1.2 Professional judgment and appointments. Contractor retains professional judgment over clinical means and methods within her license, documented competence, and the Medical Director's protocol. The parties shall coordinate appointment availability. Contractor attends accepted appointments and performs associated documentation and follow-up duties; no fixed shifts or on-site waiting periods between appointments are required by this Agreement. No minimum number of bookings, minimum earnings, or guaranteed patient volume is promised.
1.3 Resources and authority. Company provides the treatment facility, equipment, approved supplies, booking systems, and administrative support. Contractor may not bind Company, incur expenses, order products on Company's behalf, or represent that she has authority to do so except as expressly approved in writing. Other professional work is permitted subject to patient safety, accepted appointments, confidentiality, and Section 10.
## 2 Term and replacement of prior agreement
2.1 Term. This Agreement takes effect December 1, 2026, provided both parties execute it. The initial term ends April 30, 2027. It automatically renews for successive five-month terms unless either party gives written notice of non-renewal at least 30 calendar days before the current term ends. Either party may separately terminate under Section 11.
2.2 Prior agreement. As of the Effective Date, this Agreement replaces the parties' independent contractor agreement effective April 10, 2026. Rights, obligations, and earned compensation attributable to periods before December 1, 2026 remain governed by the prior agreement and any signed amendments. This Agreement does not release accrued claims or retroactively change compensation. Each party shall enter the actual signature date.
## 3 Scope of services
3.1 Authorized procedures. Subject to the current written Delegation Protocol, documented training, insurance coverage, and applicable law, services may include cosmetic botulinum toxin injections; hyaluronic acid and calcium hydroxylapatite dermal fillers, including lip fillers and Radiesse; platelet-rich plasma treatments; IPL; microneedling with or without PRP; and approved IV therapy. Biorevitalization or other additional treatments may be performed only after the Medical Director has specifically approved the procedure and products in writing and the requirements of this section have been met. A service name in this Agreement alone is not clinical authorization.
3.2 Related duties. Services include appropriate assessments, informed consent, treatment records, post-treatment instructions, clinically appropriate follow-up, and complication management and escalation. All treatments must be on Company's approved service menu and within the current Delegation Protocol attached as Exhibit A.
3.3 Restrictions. Contractor shall not independently introduce services or products, expand her scope, perform a procedure for which competency has not been documented, or treat at another location under this Agreement without prior written approval from Company and the Medical Director. Requests for off-menu services must be referred to Company and the Medical Director. All treatment decisions must remain clinically appropriate and voluntary for the patient.
## 4 Medical Director oversight and competency
4.1 Supervision. Contractor shall practice under the established Medical Director protocol and the supervision required by applicable law and that protocol. Required supervision must be available before the affected treatment begins. This Agreement does not substitute for a clinical protocol or grant prescribing, purchasing, or clinical authority beyond applicable law.
4.2 Escalation. Contractor shall promptly notify the Medical Director of complications, adverse events, contraindications, uncertainty about treatment suitability, and matters beyond her authorized competence, and shall follow applicable emergency procedures. Necessary emergency care must not be delayed for administrative approval.
4.3 Competency records. Exhibit B shall identify the procedures for which training and competency are documented and any further training or supervision required. Company-provided mentoring supplements and does not replace required physician oversight. The Medical Director must approve clinical authorization; an increase in compensation does not expand scope or remove supervision requirements.
4.4 Protocol changes. Company shall furnish the current protocol and provide notice of changes before they apply to Contractor's services, except immediate patient-safety measures. Contractor shall acknowledge receipt. A conflict with law must be resolved in favor of law; affected services shall pause pending clarification. A protocol change cannot alter compensation without a signed amendment.
## 5 Compensation and progression
5.1 Base percentage. For qualifying services personally performed on or after December 1, 2026, Company shall pay Contractor 35% of Net Collected Treatment Revenue. Company retains the balance to cover its costs and business return. Services performed before that date remain subject to the prior agreed rate, even if payment is collected later. The applicable percentage is determined by the service date unless a later signed amendment expressly provides otherwise.
5.2 Revenue definition. Net Collected Treatment Revenue means amounts actually collected and attributable to qualifying completed treatments personally performed by Contractor, after patient discounts, refunds, chargebacks, and any payment-processing fees actually borne by Company and directly attributable to those collections. Patient-paid processing surcharges are excluded from treatment revenue, and processing fees recovered through those surcharges shall not be deducted again. Product acquisition costs, supplies, rent, marketing, training, medical direction, insurance, and other overhead are not deductions from this compensation base.
5.3 Software labels and exclusions. Eligible treatment revenue includes Botox, fillers, PRP, and other authorized treatments even when the booking or billing system labels the associated injectable units or treatment materials as product sales. Excluded amounts are actual take-home retail sales, taxes, gratuities, cancellation or no-show charges, services performed by other providers, and unredeemed packages, memberships, or gift balances. No amount may be counted twice.
5.4 Prepayments and packages. Collections for a prepaid package, membership, or gift balance enter the compensation base only when the corresponding treatment is performed. Allocate the amount actually paid among included treatments under Company's consistent written allocation method, disclosed to Contractor; total allocations may not exceed the amount actually collected. A displayed list price or redeemed package value alone does not establish collected revenue. For treatment receivables collected in installments, pay the applicable share as amounts are collected.
5.5 Payment and statements. Company shall calculate and pay compensation biweekly, within five business days after each two-week period ends, with a statement showing eligible collections, deductions, the applicable percentage, compensation, and tips separately. An adjustment for a refund or chargeback is limited to Contractor's share previously paid on the affected revenue, using the original applicable rate. Company shall identify the adjustment; unrelated amounts shall not be withheld. Disputed amounts shall be addressed promptly, and undisputed earned compensation shall be paid when due. All payment and adjustment provisions are subject to applicable law.
5.6 Documentation. Contractor shall complete records as required in Section 7. If missing documentation prevents verification of a particular service, Company shall promptly identify the deficiency and may defer only that disputed service's compensation to the extent lawful. The parties shall reconcile it promptly after correction; earned compensation is not forfeited for a documentation delay.
5.7 Eligibility for review toward 40%. Beginning with December 2026, Contractor qualifies for a compensation review toward 40% after completing at least 80 Qualifying Appointments in EACH of three consecutive full calendar months. There is no minimum monthly revenue requirement. Months may not be averaged, and unused appointments do not carry forward. The earliest qualifying period is December 2026 through February 2027.
5.8 Qualifying Appointment. A Qualifying Appointment is an actually completed, paid treatment visit personally performed by Contractor and supported by the clinical and payment records, including a visit paid through redemption of a previously purchased treatment package. Multiple services, injectable units, or invoice entries within one visit count as one appointment. Touch-ups, whether paid or complimentary, do not count. Also excluded are cancellations, no-shows, consultation-only visits, follow-up-only visits, complimentary treatments, and visits performed by another provider. A patient may have multiple genuine treatment visits on different dates; encounters may not be divided artificially to increase the count. At most one qualifying visit per patient per day is counted. A visit containing a separate paid primary treatment plus a touch-up counts once for the primary treatment; the touch-up adds no count. Fully refunded visits are excluded. Counts are assigned to the treatment month, not the booking date.
5.9 Review procedure and discretion. Company shall provide a monthly count within 15 calendar days after month-end. After three qualifying months, the parties shall meet within 30 calendar days after the third month ends to review a possible increase to 40%, considering clinical performance, documentation, reliability, patient experience, purchasing assistance, and practice economics. Meeting the appointment target guarantees the review, not the increase. There is no additional revenue threshold for review eligibility. Any approved increase requires a written amendment signed by both parties stating the rate, effective date, and conditions. The 35% rate continues unless that amendment is signed. A later decline in appointments does not by itself reduce an agreed rate.
5.10 Clinical independence. Appointment targets must never influence the decision to recommend or withhold treatment, discourage appropriate touch-ups or follow-up, or encourage unnecessary visits. Compensation is for personally performed services and the responsibilities described here, not for patient referrals, prescribing a particular product, or steering patients to another provider.
5.11 Tips. Tips are separate from Net Collected Treatment Revenue and are not multiplied by the 35% or any later percentage. Contractor's tip allocation shall be documented in a separate written policy furnished to Contractor and must comply with applicable law. This Agreement does not establish a new tip split or authorize Company, an owner, manager, or supervisor to retain any tip where prohibited by law. Company shall separately account for Contractor's tips on payment statements.
## 6 Insurance
6.1 Coverage confirmation. Before Contractor provides services, Company shall confirm in writing whether Contractor is a covered provider under Company's professional liability policy, including covered procedures and applicable limits. An intention to add Contractor is not evidence of coverage. Contractor shall cooperate with credentialing and promptly report matters affecting coverage.
6.2 Separate insurance. Contractor is strongly encouraged to maintain her own professional liability policy with limits of at least $1,000,000 per claim and $3,000,000 aggregate. Company may require such separate coverage on reasonable written notice. Before services begin, the parties shall document the actual coverage arrangement, any claims-made retroactive date, responsibility for any required extended reporting coverage, and any separate insurance requirement in Exhibit C. Additional-insured status shall be obtained only when available and appropriate under the policy.
6.3 Lapse. Contractor shall not perform services lacking required coverage. The affected services are immediately suspended upon a lapse or removal from a required policy and may resume only after written coverage confirmation. Each party shall promptly notify the other of a known cancellation, material reduction, or lapse. A lapse lasting more than five business days is grounds for termination under Section 11.
## 7 Documentation and patient records
7.1 Records. Contractor shall complete accurate clinical documentation in Company's designated EMR, currently Aesthetic Record, on the day of service. Records shall include relevant history, assessment, contraindications, informed consent, treatment and product details, lot numbers, dose or volume, treatment sites, required photographs, aftercare, complications, and escalation actions. Amendments must preserve the audit trail; no record may be falsified or backdated.
7.2 Consent and privacy. Contractor shall verify procedure-specific informed consent before treatment and obtain required permission for photographs and their uses. Contractor shall use patient information only for authorized care and operations, maintain confidentiality and appropriate security, and comply with applicable privacy law. Any required business associate or other privacy agreement shall be executed separately.
7.3 Custody and access. Company is custodian of records created through the practice, subject to patient rights and applicable law. Contractor may not remove, use, or disclose those records for solicitation. Company shall provide lawful, secure access needed to complete outstanding records, respond to patient care matters, comply with professional obligations, or defend a claim, including after termination. Patient access, record transfers, and legally required retention must not be obstructed.
## 8 Regulatory compliance and professional standards
8.1 Credentials. Contractor shall maintain an active unrestricted Florida APRN license, required certifications, continuing education, BLS/CPR certification, and any additional credentials required for her actual services. Contractor represents that she is qualified for services she undertakes and shall provide supporting records. Company shall maintain the business and facility authorizations required for its operations and the necessary medical oversight arrangements.
8.2 Reporting. Contractor shall notify Company and the Medical Director promptly, and no later than 24 hours after learning of a material licensing restriction, relevant disciplinary matter, professional claim, insurance lapse, or other condition affecting safe authorized practice, subject to applicable privacy and reporting law. Patient-safety emergencies require immediate escalation. Each party remains responsible for reports legally required of that party.
8.3 Compliance. Neither party shall require an unlawful act, improper referral payment, false record, or treatment outside authorized scope. If a compensation or operational provision requires adjustment for compliance, the parties shall promptly seek a lawful written revision. This does not permit unilateral reduction of earned compensation or waiver of statutory rights.
## 9 Responsibility and indemnification
9.1 Contractor responsibility. To the extent permitted by law, Contractor shall indemnify Company and its members, managers, employees, and Medical Director against third-party claims, damages, and reasonable defense costs to the extent caused by Contractor's negligent or wrongful acts, unauthorized practice, or material breach of this Agreement. Contractor is not responsible under this clause for the portion caused by another person's negligence or misconduct. Company remains responsible for its own acts and omissions.
9.2 Claims procedure. The party seeking indemnification shall give prompt notice and reasonable cooperation. Defense and settlement shall be coordinated with applicable insurers and their policy rights. No party may settle a claim imposing an admission, nonmonetary obligation, or uninsured payment on another party without that party's written consent, not to be unreasonably withheld. Separate counsel is not automatically at Contractor's expense. Nothing in this Agreement expands insurance coverage or eliminates a legal defense.
9.3 Limits and survival. No indemnity requires an unlawful payment or shifts nonwaivable employer obligations to Contractor. These provisions survive termination for covered acts occurring during the engagement, subject to applicable limitation periods and law.
## 10 Patient relationships and confidentiality
10.1 Practice relationships. Company has an interest in its patient goodwill and confidential records. Patients retain freedom of choice and all rights under applicable law; no party owns a patient. Contractor shall not misuse Company's records or confidential information to divert business.
10.2 Patient solicitation. During the engagement and for five months after termination, Contractor shall not knowingly target for solicitation to a competing practice any Company patient whom Contractor treated or with whom she had material professional contact through Company during the 24 months before termination. This restriction does not prohibit general advertising not targeted at such patients, legally required communications or continuity of care, or accepting a patient who independently seeks care without prohibited solicitation. Its scope is limited to what applicable law permits and what is reasonably necessary to protect a legitimate business interest.
10.3 Personnel. During the engagement and for five months afterward, Contractor shall not specifically solicit Company's employees or contractors with whom she worked to end their relationship with Company for a competing engagement. General recruitment not directed at those persons is excluded.
10.4 Confidential information. Contractor shall protect nonpublic patient, financial, pricing, vendor, marketing, and business information obtained through Company and use it only for authorized purposes. Confidential information excludes information lawfully public, already lawfully known, independently developed, or lawfully obtained without restriction. Required disclosures, protected reporting to authorities, and lawful communications with legal counsel are permitted. Reasonable notice of compelled disclosure shall be given where lawful. The same protections apply to Contractor's confidential pre-existing account information disclosed to Company.
10.5 Materials and remedies. Company retains ownership of materials it creates; third-party materials and Contractor's pre-existing materials remain with their respective owners. On termination, each party shall return the other's confidential property, subject to lawful retention and secure record access under Section 7. General professional knowledge and skills are not Company property. A party may seek available legal or equitable relief for breach, subject to applicable standards and security requirements. A court may narrow restrictions only as permitted by law.
## 11 Termination and transition
11.1 Without cause. Either party may terminate on at least 14 calendar days' written notice. During notice, Contractor shall complete accepted appointments and records and cooperate in clinically appropriate handoffs, unless services are suspended for patient safety or the parties agree otherwise. No party must provide care that is unsafe or unauthorized.
11.2 Cause. Company may immediately suspend affected services or terminate for loss of required licensure or supervision, a coverage lapse described in Section 6, unauthorized practice, material patient-safety risk, fraud, material credential misrepresentation, serious privacy breach, or material violation of the protocol. For a remediable non-safety breach, Company shall identify the breach and allow five business days to cure before termination for that breach; repeated breaches after written warning need not receive another cure period. Either party may terminate for the other's material breach remaining uncured for five business days after written notice.
11.3 Final compensation and access. Company may disable routine system access at termination but shall arrange secure access needed to finish records and handoffs. Contractor shall complete outstanding records within 24 hours where practicable and cooperate on any remaining corrections. Company shall pay undisputed earned amounts on the regular payment schedule, with an initial reconciliation within 30 days. Later collections for Contractor's completed services remain payable at the applicable service-date rate and shall be reported and paid on the regular cycle as collected. Permitted refund adjustments remain subject to Section 5.
11.4 Transition. Contractor shall return Company property and reasonably cooperate with transfer of care, pending treatment plans, complications, and record requests. Company shall arrange ongoing coverage. Transition duties do not create an indefinite obligation to provide unpaid new treatment. Earned-payment, confidentiality, record, indemnification, dispute, and expressly stated post-termination provisions survive to the extent applicable.
## 12 Representations
Each party represents that it has authority to enter this Agreement and will comply with obligations applicable to it. Contractor represents that the credentials and experience she supplies are accurate, that she is not subject to a conflicting contractual restriction that prevents this engagement, and that she will notify Company of changes affecting her ability to practice. Contractor shall not undertake a treatment without the documented competence and authorization it requires. Neither party represents that a contract label, signature, or purchasing account alone establishes regulatory compliance.
## 13 Governing law and disputes
Florida law governs this Agreement. Except for urgent relief or proceedings that cannot lawfully be delayed, the parties shall first attempt good-faith resolution and then nonbinding mediation with a mutually agreed mediator in Miami-Dade County, sharing the mediator's fees equally. If unresolved within 60 days after a written mediation request, either party may pursue available remedies in state courts in Miami-Dade County or federal courts with jurisdiction there. Applicable filing deadlines are not automatically extended. The prevailing party in an action concerning this Agreement may recover reasonable attorneys' fees and costs to the extent permitted by law. Nonwaivable statutory remedies remain available.
## 14 Training and purchasing assistance
14.1 Training. Company may continue to provide or arrange training and mentoring. Contractor shall participate in training necessary for the services she elects to provide and document competency as required by the Medical Director. Neither past training nor routine mentoring creates a repayment debt under this Agreement. Any future separately funded education arrangement involving repayment must be separately agreed in writing before the expense and comply with applicable law.
14.2 Purchasing contribution. The parties acknowledge that Contractor's established supplier relationships and authorized account access have helped Company obtain favorable medical-product pricing otherwise unavailable to Company and that this assistance is a factor in the agreed 35% compensation. Contractor shall use reasonable efforts during the engagement to maintain and facilitate authorized access, communicate available pricing, and assist with approved purchases. No particular supplier discount, volume, price, or savings amount is guaranteed.
14.3 Purchase controls. Company must approve products, quantities, prices, and payment arrangements before commitments on its behalf. Company bears approved acquisition costs; Contractor need not advance funds or extend personal credit without a separate written agreement. Purchases must comply with supplier account terms, applicable licensing and product requirements, and Company's clinical approvals. Contractor shall not share credentials, transfer a restricted account, or make a purchase for Company where not authorized. Maintain invoices and available product traceability, storage, and delivery records. Disclose any personal rebate, credit, or other financial interest associated with a Company purchase before approval.
14.4 Account rights and changes. Pre-existing accounts and relationships remain with their lawful owners; this Agreement does not transfer them to Company. Contractor shall notify Company within five business days after learning of a material loss or reduction of access or discounts, and the parties shall review the effect on the arrangement within 30 days after notice. Company purchases and related credits shall be accounted for transparently under approved supplier terms. A change in access does not automatically reduce compensation, create a reimbursement debt, or authorize withholding. Any compensation change must be prospective and signed by both parties; either party retains Section 11 termination rights.
## 15 General provisions
15.1 Entire agreement and amendments. This Agreement, the actual current protocol attached as Exhibit A, completed Exhibits B and C, and any separately executed applicable privacy agreement constitute the parties' agreement for the period beginning on the Effective Date. Changes require a writing signed by both parties, except authorized clinical protocol updates under Section 4. Policies and exhibits may not unilaterally change compensation, review eligibility, or termination rights.
15.2 Notices. Notices shall be delivered personally, by tracked overnight courier, or by email with affirmative delivery confirmation or acknowledgment to the addresses on the signature page. Either party may update its notice information in writing. A notice period begins on confirmed receipt. Routine scheduling messages are not contractual termination notices unless expressly identified as such.
15.3 Assignment and severability. Company may assign to a successor acquiring the practice on written notice, provided the successor assumes this Agreement and required clinical arrangements continue. Contractor may not assign personal clinical duties without written consent and required credentialing. An unenforceable provision shall be limited or severed to the extent permitted by law without invalidating the remainder. Failure to enforce a provision is not a continuing waiver.
15.4 Other provisions. Neither party is liable for failure caused by events beyond reasonable control, excluding earned payment obligations, provided prompt notice and reasonable mitigation occur. No third-party beneficiary rights are created except express indemnification rights. Electronic signatures and counterparts are permitted. Each party acknowledges an opportunity for independent legal advice and agrees to the terms actually executed, including the requirement of a signed amendment for any increase to 40%.
## Signatures
The parties sign on the actual dates below and agree to the December 1, 2026 Effective Date stated in this Agreement.
### Company
Aventura Aesthetics Group LLC d/b/a Kami Aesthetics
Authorized signature: __________________________________________________
Printed name: Lev Polshkov    Title: ______________________________________
Date signed: __________________    Notice email: __________________________
Notice address: 2999 NE 191st Street, Suite 906, Aventura, Florida 33180
### Contractor
Valeriia Tiertyshnikova, APRN
Signature: ___________________________________________________________
Date signed: __________________    Notice email: __________________________
Current address: ______________________________________________________
_____________________________________________________________________
### Medical Director acknowledgment
The undersigned acknowledges the clinical engagement, confirms the actual Delegation Protocol identified in Exhibit A, and confirms authorization only for the procedures and supervision arrangements documented there and in Exhibit B. This acknowledgment does not make the Medical Director a party to the compensation obligations or expand a separate Medical Director agreement.
Dr. Paul M. Goldberg, MD
Signature: __________________________________    Date: __________________
@@
## Exhibit A Delegation Protocol
Attach the actual current, executed clinical protocol. This reference page is not itself a Delegation Protocol and does not authorize treatment. Affected services may not proceed until the required protocol and oversight are in place.
Protocol title: _________________________________________________________
Protocol effective date and version: _______________________________________
Number of attached pages: ______________________________________________
Supervising physician and contact details: __________________________________
_____________________________________________________________________
After-hours and emergency escalation contact: ______________________________
_____________________________________________________________________
Company acknowledgment: __________________________    Date: ____________
Contractor acknowledgment: _________________________    Date: ____________
Medical Director confirmation: ________________________    Date: ____________
## Exhibit B Procedure Competency Record
Complete a separate entry for each authorized procedure or procedure group and attach supporting training records. List botulinum toxin, applicable filler types and sites, PRP, and any other approved service separately. A commercial treatment name alone is insufficient to define scope. Add copies of this page as needed.
Procedure and approved products or treatment sites: __________________________
_____________________________________________________________________
Training date and trainer or institution: ____________________________________
Supporting competency documentation: ___________________________________
_____________________________________________________________________
Required supervision and any restrictions: __________________________________
_____________________________________________________________________
Medical Director approval: ___________________________    Date: ____________
Contractor acknowledgment: _________________________    Date: ____________
@@
## Exhibit C Insurance and Administrative Confirmations
### Insurance arrangement
Identify the coverage actually in place before services begin. Attach the applicable certificate or written insurer confirmation.
Insurer and policy number: ______________________________________________
Covered provider and covered procedures: __________________________________
_____________________________________________________________________
Policy limits: ____________________    Coverage dates: _____________________
Policy type and any retroactive date: _______________________________________
Separate Contractor policy required: ______________________________________
Responsibility for any required extended reporting coverage and cost: ______________
_____________________________________________________________________
### Payment administration
Biweekly pay-period anchor date: _________________________________________
Package allocation policy title and date supplied to Contractor: ___________________
_____________________________________________________________________
Tip allocation policy title and date supplied to Contractor: _______________________
_____________________________________________________________________
The tip policy is separate from percentage compensation and remains subject to Section 5.11 and applicable law. No tip allocation percentage is established by leaving this field blank.
### Purchasing arrangements
Approved supplier relationships or authorized account arrangements, if any: __________
_____________________________________________________________________
_____________________________________________________________________
Company purchase approval contact: ______________________________________
Invoice and rebate accounting process: ____________________________________
_____________________________________________________________________
These administrative entries do not amend the 35% rate or the review requirement of at least 80 qualifying appointments in each of three consecutive full calendar months, excluding touch-ups. No monthly revenue threshold applies.
Company signature: _________________________________    Date: ____________
Contractor signature: ________________________________    Date: ____________
'''
for raw in text.strip().split('\n'):
 line=raw.strip()
 if not line: continue
 if line=='@@': doc.add_page_break(); continue
 if line.startswith('# '): doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '): doc.add_paragraph(line[4:],'Heading 2')
 else:
  p=doc.add_paragraph(line)
  p.paragraph_format.widow_control=True
  import re
  m=re.match(r'^(\d+\.\d+ [^.]+\.) ',line)
  if m:
   p.clear(); p.add_run(m.group(1)).bold=True; p.add_run(line[len(m.group(1)):])
for root in [doc.styles.element, doc._element]:
 for node in list(root.iter(qn('w:pBdr'))):
  node.getparent().remove(node)
doc.save(out)
print(out)
print('Words',len(text.split()))
