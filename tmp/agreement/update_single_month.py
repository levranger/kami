from docx import Document
p='/Users/levpolshkov/kami/output/documents/Kami_Aesthetics_Compensation_Amendment_December_2026.docx'
d=Document(p)
changes={
'5.6 Eligibility': '5.6 Eligibility for review toward 40%. Beginning with December 2026, Contractor qualifies for a compensation review toward 40% after completing at least 80 Qualifying Appointments in any single full calendar month. There is no minimum monthly revenue requirement, and consecutive qualifying months are not required. Appointments from different months may not be combined, and unused appointments do not carry forward. The earliest qualifying month is December 2026.',
'5.8 Review': '5.8 Review procedure and discretion. Company shall provide a monthly count within 15 calendar days after month-end. After any qualifying month, the parties shall meet within 30 calendar days after that month ends to review a possible increase to 40%, considering clinical performance, documentation, reliability, patient experience, purchasing assistance, and practice economics. Meeting the appointment target guarantees the review, not the increase. There is no additional revenue threshold for review eligibility. Any approved increase requires a written amendment signed by both parties stating the rate, effective date, and conditions. The 35% rate continues unless that amendment is signed. A later decline in appointments does not by itself reduce an agreed rate.'
}
for prefix,t in changes.items():
 matches=[x for x in d.paragraphs if x.text.startswith(prefix)]
 assert len(matches)==1
 x=matches[0];x.clear();label,body=t.split('. ',1);x.add_run(label+'. ').bold=True;x.add_run(body)
t='\n'.join(x.text for x in d.paragraphs)
assert 'three consecutive' not in t and 'third month' not in t and 'February 2027' not in t
assert 'any single full calendar month' in t
assert 'Touch-ups, whether paid or complimentary, do not count.' in t
d.save(p)
print('Updated single-month eligibility and review timing')
